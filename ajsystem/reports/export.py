"""Exportación de informes a Excel (.xlsx): una fila por informe, ordenadas por consecutivo AJ-….

Seguridad: todo texto se escribe como celda de texto (nunca como fórmula, aunque empiece por «=»),
y se eliminan los caracteres de control que el formato XML de Excel no admite.
"""
from __future__ import annotations

import datetime as dt
import io
from zoneinfo import ZoneInfo

from . import schema
from .service import CLOSED, STATUS_LABELS

HEADER_FILL = "0B1020"
SIGNED_FILL = "E9F7EF"
SIGNED_FONT = "157A46"


def _local(value: dt.datetime | None, tz: ZoneInfo) -> dt.datetime | None:
    if not value:
        return None
    return value.replace(tzinfo=dt.timezone.utc).astimezone(tz).replace(tzinfo=None, microsecond=0)


def _field_value(field: schema.Field, value):
    if field.type == "checkbox":
        return "Sí" if value else "No"
    if field.type == "date" and value:
        try:
            return dt.date.fromisoformat(str(value)[:10])
        except ValueError:
            return str(value)
    return value if value not in (None, "") else None


def workbook(rows: list[dict], template_code: str, tz_name: str, truncated: bool = False) -> bytes:
    from openpyxl import Workbook  # noqa: PLC0415 (import diferido: solo al exportar)
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE  # noqa: PLC0415
    from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    tz = ZoneInfo(tz_name)
    template = schema.effective_template(template_code)
    fields = template.fields
    fixed = ["Consecutivo", "Estado del informe", "Firmado por el Director", "Firmado por", "Fecha de firma", "Técnico",
             "Usuario del técnico", "Identificación del técnico", "Creado", "Cerrado", "Cerrado por",
             "Última modificación", "Revisión", "Reaperturas", "Fotos"]
    headers = fixed + [f.label.rstrip(":") for f in fields]
    date_fmt, datetime_fmt = "dd/mm/yyyy", "dd/mm/yyyy hh:mm"

    wb = Workbook()
    ws = wb.active
    ws.title = "Informes"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    signed_fill = PatternFill("solid", fgColor=SIGNED_FILL)
    signed_font = Font(bold=True, color=SIGNED_FONT)

    for r in rows:
        data = schema.data_for_template(r["template_code"], r["data"])
        signed = r["status_code"] in CLOSED and r["signed_at"] is not None
        values = [
            r["code"], "Firmado" if signed else STATUS_LABELS.get(r["status_code"], r["status_code"]),
            "Sí" if signed else "No", r["signed_by_name"] if signed else None,
            _local(r["signed_at"], tz) if signed else None,
            r["owner_name"], r["owner_username"], r["owner_id_number"],
            _local(r["created_at"], tz), _local(r["closed_at"], tz) if r["status_code"] in CLOSED else None,
            r["closed_by_name"] if r["status_code"] in CLOSED else None, _local(r["updated_at"], tz),
            r["revision"], r["reopen_count"], int(r["photos"] or 0),
        ] + [_field_value(f, data.get(f.id)) for f in fields]
        ws.append([None] * len(values))
        row_idx = ws.max_row
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col)
            if isinstance(value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
                cell.data_type = "s"  # texto literal: «=…» nunca se interpreta como fórmula
            else:
                cell.value = value
                if isinstance(value, dt.datetime):
                    cell.number_format = datetime_fmt
                elif isinstance(value, dt.date):
                    cell.number_format = date_fmt
        if signed:
            ws.cell(row=row_idx, column=2).fill = signed_fill
            ws.cell(row=row_idx, column=2).font = signed_font
    if truncated:
        ws.append([])
        ws.append([f"Se exportaron solo los primeros {len(rows)} informes: acote los filtros para obtener el resto."])

    widths = {"Consecutivo": 20, "Estado del informe": 15, "Estado": 15, "Firmado por el Director": 12,
              "Fecha de firma": 17, "Creado": 17, "Cerrado": 17, "Última modificación": 17, "Revisión": 10,
              "Reaperturas": 12, "Fotos": 8}
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, max(14, min(40, len(h) + 4)))
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
