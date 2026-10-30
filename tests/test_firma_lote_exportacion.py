"""Firma en lote del Director, estado «Firmado», exportación a Excel, marcas normalizadas y firmas recortadas."""
import datetime as dt
import io
import json

import pytest
from conftest import create_report, fill
from PIL import Image

from ajsystem import analytics
from test_director_paneles import _active_receive, _closed_report, _profile_signature
from test_features import _jpeg, upload


def _signature_png(size=(1200, 600), box=(500, 270, 700, 330)) -> bytes:
    """Trazo pequeño en medio de un lienzo blanco enorme (lo típico de una firma escaneada)."""
    img = Image.new("RGB", size, "white")
    img.paste((20, 20, 60), box)
    out = io.BytesIO()
    img.save(out, "PNG")
    return out.getvalue()


def _batch(b, data, csrf_from="/informes/por-firmar"):
    return b.post("/informes/firmar-lote", data, csrf_from=csrf_from)


# ---------------------------------------------------------------- firma en lote
def test_director_firma_en_lote_los_seleccionados(browser, make_user, q):
    _t, _bt, r1 = _closed_report(browser, make_user, q)
    _t2, _bt2, r2 = _closed_report(browser, make_user, q)
    _t3, _bt3, r3 = _closed_report(browser, make_user, q)
    director = make_user("RESPONSABLE")
    bd = browser(director)
    _profile_signature(bd)
    page = bd.get("/informes/por-firmar").get_data(as_text=True)
    assert "Firmar seleccionados" in page and f'name="ids" value="{r1}"' in page
    before = q("SELECT id, data, revision, status_code FROM reports WHERE id IN (%s, %s) ORDER BY id", (r1, r2))
    r = _batch(bd, {"modo": "seleccionados", "ids": [str(r1), str(r2)], "op_id": "x"})
    assert r.status_code == 302
    assert _active_receive(q, r1) == [{"created_by": director["id"]}]
    assert _active_receive(q, r2) == [{"created_by": director["id"]}]
    assert _active_receive(q, r3) == []
    # Firmar no altera el contenido ni el estado; repetir la operación no duplica firmas.
    assert q("SELECT id, data, revision, status_code FROM reports WHERE id IN (%s, %s) ORDER BY id", (r1, r2)) == before
    assert _batch(bd, {"modo": "seleccionados", "ids": [str(r1)]}).status_code == 302
    assert len(q("SELECT id FROM report_attachments WHERE report_id = %s AND kind = 'FIRMA_RECIBE'", (r1,))) == 1
    events = q("SELECT detail FROM audit_log WHERE event_type = 'INFORMES_FIRMADOS_EN_LOTE' AND user_id = %s "
               "ORDER BY id", (director["id"],))
    assert [json.loads(e["detail"])["firmados"] for e in events] == [2, 0]


def test_director_firma_todos_los_pendientes_y_quedan_en_verde(browser, make_user, q):
    t, bt, r1 = _closed_report(browser, make_user, q)
    r2 = create_report(bt)
    fill(bt, q, r2, accion="cerrar")
    abierto = create_report(bt)  # en elaboración: no se firma
    bd = browser(make_user("RESPONSABLE"))
    _profile_signature(bd)
    assert _batch(bd, {"modo": "todos", "volver": "inicio"}, csrf_from="/").headers["Location"].endswith("/")
    assert len(_active_receive(q, r1)) == 1 and len(_active_receive(q, r2)) == 1
    assert _active_receive(q, abierto) == []
    assert q("SELECT COUNT(*) AS n FROM reports r WHERE r.deleted_at IS NULL AND r.status_code IN ('CERRADO','ENVIADO') "
             "AND NOT EXISTS (SELECT 1 FROM report_attachments a WHERE a.report_id = r.id "
             "AND a.kind = 'FIRMA_RECIBE' AND a.deleted_at IS NULL)")[0]["n"] == 0
    # El técnico y el Director ven el informe como «Firmado» (insignia verde) y cuenta en el resumen.
    assert 'badge-firmado' in bt.get("/informes").get_data(as_text=True)
    home = bd.get("/").get_data(as_text=True)
    assert "stat-firmado" in home and "badge-firmado" in home
    assert 'badge-firmado' in bd.get(f"/informes/{r1}").get_data(as_text=True)
    page = bd.get("/consolidacion?estado=FIRMADO&codigo=").get_data(as_text=True)
    assert f"/informes/{r1}" in page and f"/informes/{abierto}\"" not in page
    pending = bd.get("/consolidacion?estado=CERRADO").get_data(as_text=True)
    assert f"/informes/{r1}\"" not in pending


def test_firma_en_lote_solo_director_y_con_firma_de_perfil(browser, make_user, q):
    _t, bt, rid = _closed_report(browser, make_user, q)
    bj = browser(make_user("JEFE"))
    _profile_signature(bj)
    assert _batch(bj, {"modo": "todos"}, csrf_from="/").status_code == 403
    assert _batch(bt, {"modo": "todos"}, csrf_from="/").status_code == 403
    assert bt.get("/informes/por-firmar").status_code == 403
    bd = browser(make_user("RESPONSABLE"))  # sin firma en el perfil
    r = _batch(bd, {"modo": "todos"})
    assert r.status_code == 302 and _active_receive(q, rid) == []
    assert "Cargue primero su firma" in bd.get(r.headers["Location"]).get_data(as_text=True)


def test_reabrir_un_informe_firmado_lo_saca_de_firmados(browser, make_user, q):
    _t, _bt, rid = _closed_report(browser, make_user, q)
    bd = browser(make_user("RESPONSABLE"))
    _profile_signature(bd)
    _batch(bd, {"modo": "seleccionados", "ids": [str(rid)]})
    from conftest import new_op
    bd.post(f"/informes/{rid}/reabrir", {"motivo": "Corrección del número de serie", "op_id": new_op()},
            csrf_from=f"/informes/{rid}")
    assert _active_receive(q, rid) == []
    assert f"/informes/{rid}\"" not in bd.get("/consolidacion?estado=FIRMADO").get_data(as_text=True)


# ---------------------------------------------------------------- exportación a Excel
def test_exportacion_xlsx_un_informe_por_fila_por_consecutivo(browser, make_user, q):
    from openpyxl import load_workbook
    t = make_user("TECNICO")
    bt = browser(t)
    r1 = create_report(bt)
    fill(bt, q, r1, accion="cerrar", marca="=HYPERLINK(\"http://x\")", numero_serie="EXP-1")
    r2 = create_report(bt)
    fill(bt, q, r2, numero_serie="EXP-2", observaciones_habilitadas="1", observaciones="Pantalla rayada")
    bd = browser(make_user("RESPONSABLE"))
    _profile_signature(bd)
    _batch(bd, {"modo": "seleccionados", "ids": [str(r1)]})
    r = bd.get(f"/consolidacion/exportar.xlsx?tecnico={t['id']}")
    assert r.status_code == 200 and r.mimetype.endswith("spreadsheetml.sheet")
    assert r.headers["Cache-Control"] == "no-store"
    ws = load_workbook(io.BytesIO(r.data)).active
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    by_code = {row[0]: dict(zip(head, row)) for row in rows[1:]}
    codes = [row[0] for row in rows[1:]]
    assert codes == sorted(codes) and len(codes) == 2
    c1, c2 = (q("SELECT code FROM reports WHERE id = %s", (i,))[0]["code"] for i in (r1, r2))
    assert by_code[c1]["Estado del informe"] == "Firmado" and by_code[c1]["Firmado por el Director"] == "Sí"
    assert by_code[c1]["Firmado por"] and isinstance(by_code[c1]["Fecha de firma"], dt.datetime)
    assert by_code[c2]["Estado del informe"] == "En elaboración" and by_code[c2]["Observaciones adicionales"] == "Pantalla rayada"
    assert by_code[c2]["Número de serie"] == "EXP-2" and by_code[c1]["Técnico"] == f"Usuario {t['username']}"
    # Un texto que empieza por «=» se guarda como texto, nunca como fórmula.
    brand_cell = next(c for c in ws[2 + codes.index(c1)] if c.value and str(c.value).startswith("=HYPERLINK"))
    assert brand_cell.data_type == "s"
    assert q("SELECT COUNT(*) AS n FROM audit_log WHERE event_type = 'INFORMES_EXPORTADOS'")[0]["n"] >= 1
    # Los técnicos no tienen acceso a la consolidación ni a la exportación.
    assert bt.get("/consolidacion/exportar.xlsx").status_code == 403


# ---------------------------------------------------------------- tablero: marcas normalizadas
def test_marcas_normalizadas_sin_distinguir_mayusculas():
    assert analytics.normalize_brands(["HP", "Hp", " hp ", "Dell", None, "dell", "DELL", "Lenovo"]) == \
        ["HP", "HP", "HP", "DELL", None, "DELL", "DELL", "Lenovo"]
    assert analytics.normalize_brands(["Hp", "hp", "Hp"]) == ["Hp", "Hp", "Hp"]  # gana la forma más usada


def test_tablero_agrupa_marcas_y_cuenta_firmados(app, browser, make_user, q):
    t = make_user("TECNICO")
    bt = browser(t)
    for serial, brand in (("M-1", "HP"), ("M-2", "Hp"), ("M-3", "hp")):
        rid = create_report(bt)
        fill(bt, q, rid, accion="cerrar", numero_serie=serial, marca=brand)
    bd = browser(make_user("RESPONSABLE"))
    _profile_signature(bd)
    _batch(bd, {"modo": "seleccionados", "ids": [str(rid)]})
    today = dt.date.today()
    with app.app_context():
        d = analytics.executive(today - dt.timedelta(days=1), today + dt.timedelta(days=1), t["id"], "UTC")
    assert [(b["label"], b["n"]) for b in d["brands"]] == [("HP", 3)]
    assert d["kpi"]["signed"] == 1 and d["kpi"]["pending_signature"] == 2
    assert {s["code"]: s["n"] for s in d["status"]}["FIRMADO"] == 1


# ---------------------------------------------------------------- firmas sin márgenes vacíos
def test_la_firma_se_recorta_para_verse_grande(browser, make_user, q):
    _t, _bt, rid = _closed_report(browser, make_user, q)
    director = make_user("RESPONSABLE")
    bd = browser(director)
    r = upload(bd, rid, _signature_png(), kind="FIRMA_RECIBE", name="firma.png")
    assert r.status_code == 201, r.get_data(as_text=True)
    info = r.get_json()
    assert info["width"] < 300 and info["height"] < 120  # 200x60 de trazo más un margen pequeño
    # Una firma de perfil guardada antes del recorte (lienzo enorme) también se recorta al firmar.
    _profile_signature(bd)
    q("UPDATE user_signatures SET content = %s, width = 1200, height = 600 WHERE user_id = %s",
      (_signature_png(), director["id"]))
    rid2 = _closed_report(browser, make_user, q)[2]
    _batch(bd, {"modo": "seleccionados", "ids": [str(rid2)]})
    row = q("SELECT width, height FROM report_attachments WHERE report_id = %s AND kind = 'FIRMA_RECIBE' "
            "AND deleted_at IS NULL", (rid2,))[0]
    assert row["width"] < 300 and row["height"] < 120


def test_pdf_con_firmas_grandes_sigue_en_una_hoja(browser, make_user, q):
    pytest.importorskip("weasyprint")
    from weasyprint import HTML

    from ajsystem import db, documents
    from ajsystem.reports import attachments
    t = make_user("TECNICO")
    bt = browser(t)
    _profile_signature(bt)
    rid = create_report(bt)
    upload(bt, rid, _jpeg())
    fill(bt, q, rid, accion="cerrar", observaciones_habilitadas="1", observaciones="Revisión completa " * 20)
    bd = browser(make_user("RESPONSABLE"))
    _profile_signature(bd)
    _batch(bd, {"modo": "seleccionados", "ids": [str(rid)]})
    app = bd.client.application
    with app.app_context():
        report = db.query_one("SELECT r.*, o.full_name AS owner_name, o.username AS owner_username "
                              "FROM reports r JOIN users o ON o.id = r.owner_id WHERE r.id = %s", (rid,))
        report["data"] = json.loads(report["data"])
        doc = documents.build_document_data(report, [], {"full_name": "x"}, db.utcnow(), "UTC",
                                            attachments.for_document(rid))
        html = documents.render_html(doc, "aj_v1.html")
    pages = HTML(string=html, base_url=str(documents.DOC_TEMPLATES_DIR),
                 url_fetcher=documents._local_only_fetcher()).render().pages
    assert len(pages) == 2  # formato con ambas firmas en una hoja + anexo fotográfico
